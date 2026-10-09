from __future__ import annotations

import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
import python_calamine
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
INPUT_FILE = ROOT / "data" / "input" / "dummy_1c_raw.xlsx"
OUTPUT_FILE = ROOT / "data" / "output" / "dummy_cleaned_data.xlsx"

TOTAL_LABELS = {"итого", "итог", "всего", "cəmi", "cemi", "yekun", "yekun cəm"}
AMOUNT_PARENTS = {
    "Задолженность",
    "Авансы",
    "Аванс",
    "Dövrün əvvəlinə saldo",
    "Dövr üzrə dövriyyə",
    "Saldo dövrün sonuna",
}
AMOUNT_CHILDREN = {"Сумма", "Сумма УЕ", "Debet", "Kredit"}
DEBITOR_FIELDS = [
    "Плательщик",
    "Счет",
    "Офиц.номер",
    "Машина",
    "Планируемая дата оплаты",
    "Дней долга",
    "Валюта счета",
]

FONT_HEADER = Font(name="Calibri", bold=True, size=11, color="FF000000")
FONT_DATA = Font(name="Calibri", bold=False, size=11, color="FF000000")
HEADER_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BOTTOM = Border(bottom=Side(style="thin", color="000000"))
MONEY_FORMAT = "#,##0.00"
INT_FORMAT = "0"
DATE_FORMAT = "DD.MM.YYYY"


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float):
        return ""
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value).replace("\xa0", " ").replace("\n", " ").strip()


def is_empty_row(row) -> bool:
    return not any(text(value) or isinstance(value, (int, float)) and not isinstance(value, bool) for value in row)


def is_total(row) -> bool:
    return text(row[0]).lower() in TOTAL_LABELS if row else False


def find_row(data, label: str, limit: int = 15):
    for index, row in enumerate(data[:limit]):
        if any(text(value) == label for value in row):
            return index
    return None


def find_col(data, label: str, limit: int = 15):
    for row in data[:limit]:
        for index, value in enumerate(row):
            if text(value) == label:
                return index
    return None


def cell(row, index):
    if index is None or index >= len(row):
        return None
    value = row[index]
    if isinstance(value, str):
        value = value.replace("\xa0", " ").replace("\n", " ").strip()
        return value or None
    if isinstance(value, bool):
        return None
    return value


def as_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    if isinstance(value, str):
        raw = value.strip().replace(" ", "").replace(",", ".")
        if not raw:
            return None
        try:
            number = float(raw)
        except ValueError:
            return None
        return int(number) if number.is_integer() else number
    return None


def as_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
    return None


def extract_department(data) -> str:
    for row in data[:10]:
        for value in row:
            raw = text(value)
            if "Подразделение" not in raw and "Seçim" not in raw:
                continue
            match = re.search(r'"([^"]+)"', raw)
            if match:
                return match.group(1).strip()
    return ""


def parse_day_month_year(raw: str):
    day, month, year = (int(part) for part in raw.split("."))
    try:
        return date(year, month, day)
    except ValueError:
        if month == 12:
            return date(year, 12, 31)
        return date(year, month + 1, 1) - timedelta(days=1)


def extract_report_date(data, filename: str):
    for row in data[:6]:
        for value in row:
            match = re.search(r"На\s+(\d{2}\.\d{2}\.\d{4})", text(value))
            if match:
                return parse_day_month_year(match.group(1))
    match = re.search(r"(\d{2}\.\d{2}\.\d{4})", filename)
    if match:
        return parse_day_month_year(match.group(1))
    return None


def amount_columns(data, sum_row_index: int):
    parents = []
    for index in range(sum_row_index - 1, max(-1, sum_row_index - 4), -1):
        found = [
            (col, text(value))
            for col, value in enumerate(data[index])
            if text(value) in AMOUNT_PARENTS
        ]
        if found:
            parents = found
            break

    columns = []
    for col, value in enumerate(data[sum_row_index]):
        child = text(value)
        if child not in AMOUNT_CHILDREN:
            continue
        parent = ""
        for parent_col, parent_name in parents:
            if parent_col <= col:
                parent = parent_name
            else:
                break
        columns.append((col, f"{parent} {child}".strip() if parent else child))

    seen = {}
    unique = []
    for col, name in columns:
        if name in seen:
            seen[name] += 1
            name = f"{name} {seen[name]}"
        else:
            seen[name] = 1
        unique.append((col, name))
    return unique


def totals_from_row(row, columns):
    totals = {}
    for col, name in columns:
        number = as_number(cell(row, col))
        if number is not None:
            totals[name] = number
    return totals


def detect_kind(data) -> str:
    labels = set()
    for row in data[:12]:
        if row:
            labels.add(text(row[0]))
    if "Документ" in labels:
        return "creditor"
    if "Müştərilər" in labels:
        return "ehs"
    if "Клиент" in labels:
        return "debitor_grouped"
    if "Плательщик" in labels:
        return "debitor_flat"
    return "unknown"


def parse_debitor(data, report_date, department_filter: str, grouped: bool):
    payer_row = find_row(data, "Плательщик")
    sum_row = None
    for index, row in enumerate(data[:15]):
        labels = {text(value) for value in row}
        if "Сумма" in labels and "Сумма УЕ" in labels:
            sum_row = index
            break
    if payer_row is None or sum_row is None:
        raise ValueError("Debitor başlığı tapılmadı")

    amount_cols = amount_columns(data, sum_row)
    field_cols = {label: find_col(data, label) for label in DEBITOR_FIELDS}
    start = max(payer_row, sum_row) + 1

    def is_detail(row):
        return cell(row, field_cols["Счет"]) is not None or cell(row, field_cols["Валюта счета"]) is not None

    headers = ["Hesabat tarixi", "Подразделение"]
    if grouped:
        headers.append("Клиент")
    headers.extend(DEBITOR_FIELDS)
    headers.extend(name for _, name in amount_cols)

    itogo = {}
    for row in data[start:]:
        if is_total(row):
            itogo = totals_from_row(row, amount_cols)

    if not grouped:
        rows = []
        for row in data[start:]:
            if is_empty_row(row) or is_total(row):
                continue
            record = [report_date, department_filter or None]
            record.extend(cell(row, field_cols[label]) for label in DEBITOR_FIELDS)
            record.extend(as_number(cell(row, col)) for col, _ in amount_cols)
            rows.append(record)
        return headers, rows, itogo, 0

    rows = []
    department = None
    seen_department = False
    current_client = None
    pending = None
    pending_details = 0
    kept_groups = 0

    def flush():
        nonlocal pending_details, kept_groups
        if pending is not None and pending_details == 0 and any(value is not None for value in pending[1]):
            record = [report_date, department, pending[0]]
            record.extend([None] * len(DEBITOR_FIELDS))
            record.extend(pending[1])
            rows.append(record)
            kept_groups += 1
        pending_details = 0

    for row in data[start:]:
        if is_empty_row(row) or is_total(row):
            continue
        if is_detail(row):
            pending_details += 1
            record = [report_date, department, current_client]
            record.extend(cell(row, field_cols[label]) for label in DEBITOR_FIELDS)
            record.extend(as_number(cell(row, col)) for col, _ in amount_cols)
            rows.append(record)
            continue
        name = text(row[0]) or None
        if not seen_department and name:
            department = name
            seen_department = True
            continue
        flush()
        current_client = name
        pending = (name, [as_number(cell(row, col)) for col, _ in amount_cols])
        pending_details = 0
        seen_department = True
    flush()
    if department is None:
        department = department_filter or None
        for record in rows:
            record[1] = department
    return headers, rows, itogo, kept_groups


def parse_creditor(data, report_date, department_filter: str):
    doc_row = find_row(data, "Документ")
    sum_row = None
    for index, row in enumerate(data[:15]):
        labels = {text(value) for value in row}
        if "Сумма" in labels and "Сумма УЕ" in labels:
            sum_row = index
            break
    if doc_row is None or sum_row is None:
        raise ValueError("Kreditor başlığı tapılmadı")

    amount_cols = amount_columns(data, sum_row)
    invoice_col = find_col(data, "№ счета поставщика")
    currency_col = find_col(data, "Валюта")
    start = doc_row + 1

    def is_detail(row):
        return cell(row, currency_col) is not None

    headers = [
        "Hesabat tarixi",
        "Подразделение",
        "Поставщик",
        "Документ",
        "№ счета поставщика",
        "Валюта",
    ]
    headers.extend(name for _, name in amount_cols)

    rows = []
    itogo = {}
    department = None
    supplier = None
    seen_department = False
    pending = None
    pending_details = 0
    kept_groups = 0

    def flush():
        nonlocal pending_details, kept_groups
        if pending is not None and pending_details == 0 and any(value is not None for value in pending[1]):
            record = [report_date, department, pending[0], None, None, None]
            record.extend(pending[1])
            rows.append(record)
            kept_groups += 1
        pending_details = 0

    for row in data[start:]:
        if is_empty_row(row):
            continue
        if is_total(row):
            itogo = totals_from_row(row, amount_cols)
            continue
        if is_detail(row):
            pending_details += 1
            record = [
                report_date,
                department,
                supplier,
                cell(row, 0),
                cell(row, invoice_col),
                cell(row, currency_col),
            ]
            record.extend(as_number(cell(row, col)) for col, _ in amount_cols)
            rows.append(record)
            continue
        name = text(row[0]) or None
        if not seen_department and name:
            department = name
            seen_department = True
            continue
        flush()
        supplier = name
        pending = (name, [as_number(cell(row, col)) for col, _ in amount_cols])
        pending_details = 0
        seen_department = True
    flush()
    if department is None:
        department = department_filter or None
        for record in rows:
            record[1] = department
    return headers, rows, itogo, kept_groups


def parse_ehs(data, report_date, department_filter: str, filename: str):
    customer_row = find_row(data, "Müştərilər")
    if customer_row is None:
        raise ValueError("Account header was not found")
    amount_cols = amount_columns(data, customer_row)
    bolme = department_filter or None
    headers = ["Hesabat tarixi", "Bölmə", "Hesab", "Müştərilər"]
    headers.extend(name for _, name in amount_cols)

    account = None
    title = text(data[0][0]) if data and data[0] else ""
    title_account = re.match(r"(\d+)", title)
    if title_account:
        account = title_account.group(1)

    rows = []
    itogo = {}
    for row in data[customer_row + 1 :]:
        if is_empty_row(row):
            continue
        if is_total(row):
            itogo = totals_from_row(row, amount_cols)
            continue
        name = text(row[0]) or None
        if name and re.fullmatch(r"\d+(\.\d+)*", name):
            account = name
            continue
        record = [report_date, bolme, account, name]
        record.extend(as_number(cell(row, col)) for col, _ in amount_cols)
        rows.append(record)
    return headers, rows, itogo, 0


def column_kind(header: str) -> str:
    lowered = header.lower()
    if "tarix" in lowered or "дата" in lowered:
        return "date"
    if "Дней" in header:
        return "int"
    if any(token in header for token in ("Сумма", "Debet", "Kredit")):
        return "money"
    return "text"


def normalize_record(headers, record):
    normalized = []
    for header, value in zip(headers, record):
        kind = column_kind(header)
        if kind == "date":
            parsed = as_date(value)
            normalized.append(parsed if parsed is not None else (None if value in ("", None) else value))
        elif kind == "money":
            normalized.append(as_number(value) if not isinstance(value, (int, float)) or isinstance(value, bool) else value)
        elif kind == "int":
            number = as_number(value)
            normalized.append(number)
        else:
            if isinstance(value, str):
                value = value.strip() or None
            normalized.append(value)
    return normalized


def column_sum(headers, rows, name: str) -> float:
    if name not in headers:
        return 0.0
    index = headers.index(name)
    total = 0.0
    for record in rows:
        value = record[index]
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            total += float(value)
    return total


def compare_totals(headers, rows, itogo, kind: str):
    mismatches = []
    if not itogo:
        return mismatches

    if kind == "ehs":
        pairs = [
            ("Dövrün əvvəlinə saldo Debet", "Dövrün əvvəlinə saldo Kredit"),
            ("Saldo dövrün sonuna Debet", "Saldo dövrün sonuna Kredit"),
        ]
        for debet_name, kredit_name in pairs:
            expected = float(itogo.get(debet_name, 0) or 0) - float(itogo.get(kredit_name, 0) or 0)
            actual = column_sum(headers, rows, debet_name) - column_sum(headers, rows, kredit_name)
            if round(actual - expected, 2) != 0:
                mismatches.append((f"{debet_name} - {kredit_name}", expected, round(actual, 2), round(actual - expected, 2)))
        for name, expected in itogo.items():
            if "dövriyyə" not in name.lower():
                continue
            actual = column_sum(headers, rows, name)
            if round(actual - float(expected), 2) != 0:
                mismatches.append((name, expected, round(actual, 2), round(actual - float(expected), 2)))
        return mismatches

    for name, expected in itogo.items():
        actual = column_sum(headers, rows, name)
        if round(actual - float(expected), 2) != 0:
            mismatches.append((name, expected, round(actual, 2), round(actual - float(expected), 2)))
    return mismatches


def write_workbook(path: Path, frame: pd.DataFrame):
    headers = list(frame.columns)
    records = frame.to_dict(orient="records")
    rows = [
        [record[header] if pd.notna(record[header]) else None for header in headers]
        for record in records
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    for col, header in enumerate(headers, start=1):
        cell_obj = sheet.cell(1, col, header)
        cell_obj.font = FONT_HEADER
        cell_obj.alignment = HEADER_ALIGN
        cell_obj.border = THIN_BOTTOM

    kinds = [column_kind(header) for header in headers]
    for row_index, record in enumerate(rows, start=2):
        for col, (value, kind) in enumerate(zip(record, kinds), start=1):
            cell_obj = sheet.cell(row_index, col, value)
            cell_obj.font = FONT_DATA
            if kind == "money" and isinstance(value, (int, float)):
                cell_obj.number_format = MONEY_FORMAT
            elif kind == "int" and isinstance(value, (int, float)):
                cell_obj.number_format = INT_FORMAT
            elif kind == "date" and isinstance(value, date):
                cell_obj.number_format = DATE_FORMAT

    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 32

    for col, header in enumerate(headers, start=1):
        length = len(str(header))
        for record in rows[:200]:
            value = record[col - 1]
            if value is not None:
                length = max(length, min(len(str(value)), 60))
        sheet.column_dimensions[get_column_letter(col)].width = min(max(length + 3, 14), 48)

    workbook.save(path)


def process_file(path: Path):
    workbook = python_calamine.CalamineWorkbook.from_path(str(path))
    data = workbook.get_sheet_by_index(0).to_python()
    kind = detect_kind(data)
    report_date = extract_report_date(data, path.name)
    department = extract_department(data)

    if kind == "creditor":
        headers, rows, itogo, kept = parse_creditor(data, report_date, department)
    elif kind == "ehs":
        headers, rows, itogo, kept = parse_ehs(data, report_date, department, path.name)
    elif kind == "debitor_grouped":
        headers, rows, itogo, kept = parse_debitor(data, report_date, department, grouped=True)
    elif kind == "debitor_flat":
        headers, rows, itogo, kept = parse_debitor(data, report_date, department, grouped=False)
    else:
        raise ValueError("Unrecognized report layout")

    rows = [normalize_record(headers, record) for record in rows]
    mismatches = compare_totals(headers, rows, itogo, kind)
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    frame = pd.DataFrame(rows, columns=headers)
    write_workbook(OUTPUT_FILE, frame)
    return {
        "file": path.name,
        "kind": kind,
        "source_rows": len(data),
        "rows": len(rows),
        "kept_groups": kept,
        "mismatches": mismatches,
        "output": OUTPUT_FILE.name,
    }


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if not INPUT_FILE.exists():
        print(f"Input file was not found: {INPUT_FILE}")
        return 1
    result = process_file(INPUT_FILE)
    print(f"{result['kind']}: {result['rows']} rows -> {OUTPUT_FILE}")
    if result["mismatches"]:
        for name, expected, actual, diff in result["mismatches"]:
            print(f"{name}: report {expected} | sum {actual} | diff {diff}")
        return 1
    print("Totals match the report total row.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
